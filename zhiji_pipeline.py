"""知迹：ASSISTments 下一次作答预测与轻量 ENA 分析。

设计原则：
1. 按学生与时间重建轨迹；2. 按学生划分训练/验证/测试集；
3. 所有行为和情感特征均先 shift(1)，避免看到待预测作答；
4. 同时提供 XGBoost 基线、GRU 序列模型与 ENA 风格共现网络。
"""

from __future__ import annotations

import argparse
import json
import random
from collections import defaultdict
from pathlib import Path

import joblib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
import torch
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.metrics import accuracy_score, f1_score, log_loss, recall_score, roc_auc_score
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from torch import nn
from torch.utils.data import DataLoader, Dataset
try:
    from xgboost import XGBClassifier
    HAS_XGBOOST = True
except ImportError:
    from sklearn.ensemble import RandomForestClassifier
    HAS_XGBOOST = False

SEED = 42
EMOTION_COLUMNS = {
    "Average_confidence(FRUSTRATED)": "frustrated",
    "Average_confidence(CONFUSED)": "confused",
    "Average_confidence(CONCENTRATING)": "concentrating",
    "Average_confidence(BORED)": "bored",
}
KEEP_COLUMNS = [
    "problem_log_id", "user_id", "problem_id", "skill", "skill_id",
    "assignment_id", "start_time", "correct", "hint_count", "attempt_count",
    "ms_first_response", *EMOTION_COLUMNS,
]
HISTORY_BASE = [
    "correct", "hint_count", "attempt_count", "response_seconds",
    "frustrated", "confused", "concentrating", "bored",
]


def seed_everything(seed: int = SEED) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)


def read_dataset(path: Path, max_rows: int | None = None) -> pd.DataFrame:
    """读取 CSV 或大型 XLSX；XLSX 使用只读流式模式。"""
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, usecols=lambda c: c in KEEP_COLUMNS, nrows=max_rows)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 CSV、XLSX 或 XLSM")

    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(rows)]
    wanted = [(i, name) for i, name in enumerate(header) if name in KEEP_COLUMNS]
    records = []
    for n, row in enumerate(rows, start=1):
        if max_rows is not None and n > max_rows:
            break
        if not any(v is not None for v in row):
            continue
        records.append({name: row[i] for i, name in wanted})
    wb.close()
    return pd.DataFrame.from_records(records)


def prepare_interactions(raw: pd.DataFrame) -> pd.DataFrame:
    df = raw.rename(columns=EMOTION_COLUMNS).copy()
    required = {"user_id", "problem_log_id", "start_time", "correct"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"数据缺少必要字段：{sorted(missing)}")

    numeric = ["correct", "hint_count", "attempt_count", "ms_first_response",
               "frustrated", "confused", "concentrating", "bored"]
    for col in numeric:
        if col not in df:
            df[col] = np.nan
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["start_time"] = pd.to_datetime(df["start_time"], errors="coerce")
    df["response_seconds"] = df["ms_first_response"].clip(lower=0) / 1000.0
    df["skill"] = df.get("skill", pd.Series(index=df.index, dtype=object)).fillna("Unknown").astype(str)

    # 同一次交互若因多个知识点重复，仅保留一条；技能名合并便于追溯。
    aggregations = {c: "first" for c in df.columns if c not in {"problem_log_id", "skill"}}
    aggregations["skill"] = lambda s: " | ".join(sorted(set(s.dropna().astype(str))))
    df = df.groupby("problem_log_id", as_index=False, sort=False).agg(aggregations)
    df = df.dropna(subset=["user_id", "start_time", "correct"])
    df = df[df["correct"].isin([0, 1])]
    return df.sort_values(["user_id", "start_time", "problem_log_id"]).reset_index(drop=True)


def add_history_features(df: pd.DataFrame, window: int = 5) -> pd.DataFrame:
    out = df.copy()
    group = out.groupby("user_id", sort=False)
    for col in HISTORY_BASE:
        out[f"prev_{col}"] = group[col].shift(1)
    out["history_count"] = group.cumcount()
    out["rolling_correct_rate"] = group["correct"].transform(
        lambda s: s.shift(1).rolling(window, min_periods=1).mean()
    )
    out["rolling_hint_rate"] = group["hint_count"].transform(
        lambda s: s.shift(1).gt(0).rolling(window, min_periods=1).mean()
    )
    out["rolling_error_count"] = group["correct"].transform(
        lambda s: s.shift(1).eq(0).rolling(window, min_periods=1).sum()
    )
    # 至少有一次历史交互时才能预测“下一次”。
    return out[out["history_count"] > 0].reset_index(drop=True)


def split_students(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    first = GroupShuffleSplit(n_splits=1, test_size=0.20, random_state=SEED)
    train_val_idx, test_idx = next(first.split(df, groups=df["user_id"]))
    train_val, test = df.iloc[train_val_idx], df.iloc[test_idx]
    second = GroupShuffleSplit(n_splits=1, test_size=0.20, random_state=SEED)
    train_idx, val_idx = next(second.split(train_val, groups=train_val["user_id"]))
    return train_val.iloc[train_idx].copy(), train_val.iloc[val_idx].copy(), test.copy()


def metrics(y: np.ndarray, prob: np.ndarray) -> dict[str, float]:
    pred = (prob >= 0.5).astype(int)
    result = {
        "accuracy": accuracy_score(y, pred),
        "macro_f1": f1_score(y, pred, average="macro"),
        "incorrect_recall": recall_score(y, pred, pos_label=0),
        "log_loss": log_loss(y, np.c_[1 - prob, prob], labels=[0, 1]),
    }
    result["auc"] = roc_auc_score(y, prob) if len(np.unique(y)) == 2 else float("nan")
    return {k: float(v) for k, v in result.items()}


def train_xgboost(train: pd.DataFrame, val: pd.DataFrame, test: pd.DataFrame, out: Path) -> dict:
    numeric = [f"prev_{c}" for c in HISTORY_BASE] + [
        "history_count", "rolling_correct_rate", "rolling_hint_rate", "rolling_error_count"
    ]
    categorical = ["skill"]
    transformer = ColumnTransformer([
        ("num", Pipeline([("impute", SimpleImputer(strategy="median")),
                          ("scale", StandardScaler())]), numeric),
        ("cat", Pipeline([("impute", SimpleImputer(strategy="most_frequent")),
                          ("onehot", OneHotEncoder(handle_unknown="ignore", min_frequency=5))]), categorical),
    ])
    if HAS_XGBOOST:
        model = XGBClassifier(
            n_estimators=350, max_depth=5, learning_rate=0.05, subsample=0.85,
            colsample_bytree=0.85, eval_metric="logloss", n_jobs=1, random_state=SEED,
        )
    else:
        model = RandomForestClassifier(
            n_estimators=250, max_depth=12, min_samples_leaf=5,
            class_weight="balanced_subsample", n_jobs=1, random_state=SEED,
        )
    pipe = Pipeline([("features", transformer), ("model", model)])
    pipe.fit(train[numeric + categorical], train["correct"].astype(int))
    joblib.dump(pipe, out / "xgboost_pipeline.joblib")
    return {
        "validation": metrics(val["correct"].to_numpy(), pipe.predict_proba(val[numeric + categorical])[:, 1]),
        "test": metrics(test["correct"].to_numpy(), pipe.predict_proba(test[numeric + categorical])[:, 1]),
    }


class SequenceDataset(Dataset):
    def __init__(self, frame: pd.DataFrame, skill_map: dict[str, int], length: int = 20):
        self.samples = []
        cols = HISTORY_BASE
        for _, part in frame.groupby("user_id", sort=False):
            part = part.sort_values("start_time")
            skills = part["skill"].map(skill_map).fillna(0).astype(int).to_numpy()
            values = part[cols].fillna(0).astype(float).to_numpy(np.float32)
            labels = part["correct"].astype(int).to_numpy()
            # 输入截止到 t-1，目标为 t。
            for t in range(1, len(part)):
                begin = max(0, t - length)
                self.samples.append((skills[begin:t], values[begin:t], labels[t]))
        self.length = length

    def __len__(self): return len(self.samples)

    def __getitem__(self, index):
        skills, values, label = self.samples[index]
        pad = self.length - len(skills)
        return (
            torch.tensor(np.pad(skills, (pad, 0))),
            torch.tensor(np.pad(values, ((pad, 0), (0, 0)))),
            torch.tensor(len(skills) - 1),
            torch.tensor(label, dtype=torch.float32),
        )


class GRUPredictor(nn.Module):
    def __init__(self, n_skills: int, n_cont: int, hidden: int = 64):
        super().__init__()
        self.embedding = nn.Embedding(n_skills + 1, 24, padding_idx=0)
        self.gru = nn.GRU(24 + n_cont, hidden, batch_first=True)
        self.head = nn.Sequential(nn.Linear(hidden, 32), nn.ReLU(), nn.Dropout(0.2), nn.Linear(32, 1))

    def forward(self, skills, values, last_index):
        sequence, _ = self.gru(torch.cat([self.embedding(skills), values], dim=-1))
        row = torch.arange(sequence.size(0), device=sequence.device)
        # 左侧 padding，因此最后一个真实事件恒在最后一列。
        return self.head(sequence[row, -1]).squeeze(-1)


def train_gru(train: pd.DataFrame, val: pd.DataFrame, test: pd.DataFrame, out: Path,
              epochs: int = 12, length: int = 20) -> dict:
    skills = train["skill"].value_counts()
    skill_map = {s: i + 1 for i, s in enumerate(skills.index)}
    datasets = [SequenceDataset(x, skill_map, length) for x in (train, val, test)]
    loaders = [DataLoader(x, batch_size=256, shuffle=(i == 0)) for i, x in enumerate(datasets)]
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model = GRUPredictor(len(skill_map), len(HISTORY_BASE)).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=1e-3)
    loss_fn = nn.BCEWithLogitsLoss()
    best, best_loss, patience = None, float("inf"), 0

    def evaluate(loader):
        model.eval(); ys, ps, losses = [], [], []
        with torch.no_grad():
            for sk, vals, idx, y in loader:
                logits = model(sk.to(device), vals.to(device), idx.to(device))
                losses.append(loss_fn(logits, y.to(device)).item())
                ys.extend(y.numpy()); ps.extend(torch.sigmoid(logits).cpu().numpy())
        return np.asarray(ys), np.asarray(ps), float(np.mean(losses))

    for _ in range(epochs):
        model.train()
        for sk, vals, idx, y in loaders[0]:
            optimizer.zero_grad()
            loss = loss_fn(model(sk.to(device), vals.to(device), idx.to(device)), y.to(device))
            loss.backward(); optimizer.step()
        _, _, val_loss = evaluate(loaders[1])
        if val_loss < best_loss:
            best_loss, patience = val_loss, 0
            best = {k: v.detach().cpu().clone() for k, v in model.state_dict().items()}
        else:
            patience += 1
            if patience >= 3: break
    model.load_state_dict(best)
    torch.save({"state_dict": best, "skill_map": skill_map, "history_columns": HISTORY_BASE}, out / "gru_model.pt")
    vy, vp, _ = evaluate(loaders[1]); ty, tp, _ = evaluate(loaders[2])
    return {"validation": metrics(vy, vp), "test": metrics(ty, tp)}


def ena_codes(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    slow_cut = out["response_seconds"].quantile(0.75)
    out["Incorrect"] = out["correct"].eq(0)
    out["HintSeeking"] = out["hint_count"].fillna(0).gt(0)
    out["RepeatedAttempt"] = out["attempt_count"].fillna(1).gt(1)
    out["SlowResponse"] = out["response_seconds"].gt(slow_cut)
    out["ConfusionHigh"] = out["confused"].gt(0.5)
    out["FrustrationHigh"] = out["frustrated"].gt(0.5)
    out["BoredomHigh"] = out["bored"].gt(0.5)
    out["ConcentrationHigh"] = out["concentrating"].gt(0.5)
    return out


def build_ena(df: pd.DataFrame, out: Path, window: int = 5) -> None:
    codes = ["Incorrect", "HintSeeking", "RepeatedAttempt", "SlowResponse",
             "ConfusionHigh", "FrustrationHigh", "BoredomHigh", "ConcentrationHigh"]
    frame = ena_codes(df)
    edges = {"stable": defaultdict(float), "difficulty": defaultdict(float)}
    units = defaultdict(lambda: defaultdict(set))
    for user, part in frame.groupby("user_id", sort=False):
        part = part.sort_values("start_time").reset_index(drop=True)
        future_errors = sum(part["correct"].shift(-step).eq(0).astype(int) for step in (1, 2, 3))
        has_three_future = part["correct"].shift(-3).notna()
        for end in range(len(part)):
            start = max(0, end - window + 1)
            active = [c for c in codes if part.loc[start:end, c].any()]
            if not has_three_future.iloc[end]:
                continue
            group = "difficulty" if future_errors.iloc[end] >= 2 else "stable"
            for i, a in enumerate(active):
                for b in active[i + 1:]: units[group][user].add(tuple(sorted((a, b))))
    for group in edges:
        n_users = max(len(units[group]), 1)
        for pairs in units[group].values():
            for pair in pairs: edges[group][pair] += 1 / n_users
        pd.DataFrame([(a, b, w) for (a, b), w in edges[group].items()],
                     columns=["source", "target", "weight"]).to_csv(out / f"ena_{group}_edges.csv", index=False)

    graph = nx.Graph()
    for pair in set(edges["stable"]) | set(edges["difficulty"]):
        graph.add_edge(*pair, weight=edges["difficulty"].get(pair, 0) - edges["stable"].get(pair, 0))
    pos = nx.circular_layout(graph)
    widths = [1 + 7 * abs(d["weight"]) for *_, d in graph.edges(data=True)]
    colors = ["#D1495B" if d["weight"] > 0 else "#2E86AB" for *_, d in graph.edges(data=True)]
    plt.figure(figsize=(10, 8)); nx.draw_networkx(graph, pos, width=widths, edge_color=colors,
        node_color="#F4F7FA", edgecolors="#334", node_size=2300, font_size=9)
    plt.title("ENA difference: difficulty (red) minus stable (blue)"); plt.axis("off"); plt.tight_layout()
    plt.savefig(out / "ena_difference.png", dpi=180); plt.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, default=Path("2012-2013-data-with-predictions-4-final.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("results"))
    parser.add_argument("--max-rows", type=int, default=300_000,
                        help="先用子样本调试；传 0 表示读取全量")
    parser.add_argument("--epochs", type=int, default=12)
    args = parser.parse_args(); seed_everything(); args.output.mkdir(parents=True, exist_ok=True)
    raw = read_dataset(args.data, None if args.max_rows == 0 else args.max_rows)
    interactions = prepare_interactions(raw)
    featured = add_history_features(interactions)
    train, val, test = split_students(featured)
    summary = {
        "rows_raw": len(raw), "interactions": len(interactions), "students": int(interactions.user_id.nunique()),
        "split_students": {"train": int(train.user_id.nunique()), "validation": int(val.user_id.nunique()),
                           "test": int(test.user_id.nunique())},
    }
    summary["xgboost"] = train_xgboost(train, val, test, args.output)
    summary["gru"] = train_gru(train, val, test, args.output, args.epochs)
    build_ena(interactions, args.output)
    (args.output / "metrics.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
